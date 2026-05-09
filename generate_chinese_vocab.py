from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Page margins
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

def set_font(run, bold=False, size=12, color=None):
    run.font.name = '微软雅黑'
    run.font.bold = bold
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, bold=True, size=18, color=(0x1F, 0x49, 0x7D))
    return p

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, bold=False, size=11, color=(0x70, 0x70, 0x70))
    return p

def add_section_heading(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, bold=True, size=13, color=(0xFF, 0xFF, 0xFF))
    # Shade the paragraph background
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '2E75B6')
    pPr.append(shd)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.3)
    return p

def add_vocab_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for run in hdr_cells[i].paragraphs[0].runs:
            set_font(run, bold=True, size=11, color=(0xFF, 0xFF, 0xFF))
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '2E75B6')
        tcPr.append(shd)

    # Data rows
    for idx, row_data in enumerate(rows):
        row_cells = table.add_row().cells
        fill = 'DEEAF1' if idx % 2 == 0 else 'FFFFFF'
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = cell_text
            for run in row_cells[i].paragraphs[0].runs:
                set_font(run, bold=False, size=11)
            row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            tc = row_cells[i]._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), fill)
            tcPr.append(shd)

    doc.add_paragraph()

def add_tip_box(tips):
    p = doc.add_paragraph()
    run = p.add_run('📌 建议学习方法：')
    set_font(run, bold=True, size=12, color=(0x1F, 0x49, 0x7D))
    for tip in tips:
        tp = doc.add_paragraph(style='List Bullet')
        run = tp.add_run(tip)
        set_font(run, bold=False, size=11)

# ── TITLE ──────────────────────────────────────────────
add_title('新加坡小六华文常用词汇手册')
add_subtitle('Lower Chinese | Primary 6 | 小学六年级')
doc.add_paragraph()

# ── 1. 人物与家庭 ───────────────────────────────────────
add_section_heading('一、人物与家庭')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('爸爸', 'bà ba', 'father'),
        ('妈妈', 'mā ma', 'mother'),
        ('哥哥', 'gē ge', 'older brother'),
        ('姐姐', 'jiě jie', 'older sister'),
        ('弟弟', 'dì di', 'younger brother'),
        ('妹妹', 'mèi mei', 'younger sister'),
        ('祖父', 'zǔ fù', 'grandfather (paternal)'),
        ('祖母', 'zǔ mǔ', 'grandmother (paternal)'),
        ('外公', 'wài gōng', 'grandfather (maternal)'),
        ('外婆', 'wài pó', 'grandmother (maternal)'),
        ('朋友', 'péng yǒu', 'friend'),
        ('同学', 'tóng xué', 'classmate'),
        ('老师', 'lǎo shī', 'teacher'),
        ('邻居', 'lín jū', 'neighbour'),
    ]
)

# ── 2. 学校生活 ─────────────────────────────────────────
add_section_heading('二、学校生活')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('学校', 'xué xiào', 'school'),
        ('课室', 'kè shì', 'classroom'),
        ('图书馆', 'tú shū guǎn', 'library'),
        ('食堂', 'shí táng', 'canteen'),
        ('操场', 'cāo chǎng', 'school field'),
        ('功课', 'gōng kè', 'homework'),
        ('考试', 'kǎo shì', 'exam'),
        ('成绩', 'chéng jì', 'results'),
        ('课本', 'kè běn', 'textbook'),
        ('铅笔', 'qiān bǐ', 'pencil'),
        ('橡皮', 'xiàng pí', 'eraser'),
        ('尺子', 'chǐ zi', 'ruler'),
    ]
)

# ── 3. 日常动词 ─────────────────────────────────────────
add_section_heading('三、日常动词')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('吃', 'chī', 'eat'),
        ('喝', 'hē', 'drink'),
        ('走', 'zǒu', 'walk'),
        ('跑', 'pǎo', 'run'),
        ('看', 'kàn', 'look / watch'),
        ('听', 'tīng', 'listen'),
        ('说', 'shuō', 'speak'),
        ('写', 'xiě', 'write'),
        ('读', 'dú', 'read'),
        ('做', 'zuò', 'do / make'),
        ('买', 'mǎi', 'buy'),
        ('卖', 'mài', 'sell'),
        ('帮助', 'bāng zhù', 'help'),
        ('回答', 'huí dá', 'answer'),
        ('提问', 'tí wèn', 'ask a question'),
        ('整理', 'zhěng lǐ', 'tidy up'),
        ('打扫', 'dǎ sǎo', 'clean / sweep'),
        ('准备', 'zhǔn bèi', 'prepare'),
    ]
)

# ── 4. 形容词 ───────────────────────────────────────────
add_section_heading('四、形容词（描述词）')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('高兴', 'gāo xìng', 'happy'),
        ('难过', 'nán guò', 'sad'),
        ('生气', 'shēng qì', 'angry'),
        ('害怕', 'hài pà', 'afraid'),
        ('开心', 'kāi xīn', 'joyful'),
        ('聪明', 'cōng míng', 'smart'),
        ('勤劳', 'qín láo', 'hardworking'),
        ('懒惰', 'lǎn duò', 'lazy'),
        ('勇敢', 'yǒng gǎn', 'brave'),
        ('诚实', 'chéng shí', 'honest'),
        ('美丽', 'měi lì', 'beautiful'),
        ('干净', 'gān jìng', 'clean'),
        ('整齐', 'zhěng qí', 'tidy'),
        ('热闹', 'rè nào', 'lively / bustling'),
    ]
)

# ── 5. 时间词 ───────────────────────────────────────────
add_section_heading('五、时间词')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('今天', 'jīn tiān', 'today'),
        ('昨天', 'zuó tiān', 'yesterday'),
        ('明天', 'míng tiān', 'tomorrow'),
        ('早上', 'zǎo shàng', 'morning'),
        ('下午', 'xià wǔ', 'afternoon'),
        ('晚上', 'wǎn shàng', 'evening / night'),
        ('上个星期', 'shàng gè xīng qī', 'last week'),
        ('这个星期', 'zhè gè xīng qī', 'this week'),
        ('下个星期', 'xià gè xīng qī', 'next week'),
        ('月', 'yuè', 'month'),
        ('年', 'nián', 'year'),
    ]
)

# ── 6. 饮食与食物 ───────────────────────────────────────
add_section_heading('六、饮食与食物')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('米饭', 'mǐ fàn', 'rice'),
        ('面条', 'miàn tiáo', 'noodles'),
        ('面包', 'miàn bāo', 'bread'),
        ('蔬菜', 'shū cài', 'vegetables'),
        ('水果', 'shuǐ guǒ', 'fruits'),
        ('鱼', 'yú', 'fish'),
        ('鸡肉', 'jī ròu', 'chicken'),
        ('猪肉', 'zhū ròu', 'pork'),
        ('牛奶', 'niú nǎi', 'milk'),
        ('饮料', 'yǐn liào', 'drinks / beverages'),
    ]
)

# ── 7. 交通出行 ─────────────────────────────────────────
add_section_heading('七、交通出行')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('巴士', 'bā shì', 'bus'),
        ('地铁', 'dì tiě', 'MRT'),
        ('德士', 'dé shì', 'taxi'),
        ('脚踏车', 'jiǎo tà chē', 'bicycle'),
        ('汽车', 'qì chē', 'car'),
        ('飞机', 'fēi jī', 'airplane'),
        ('路口', 'lù kǒu', 'junction / intersection'),
        ('马路', 'mǎ lù', 'road'),
    ]
)

# ── 8. 自然与环境 ───────────────────────────────────────
add_section_heading('八、自然与环境')
add_vocab_table(
    ['词语', '拼音', '英文'],
    [
        ('天气', 'tiān qì', 'weather'),
        ('晴天', 'qíng tiān', 'sunny day'),
        ('雨天', 'yǔ tiān', 'rainy day'),
        ('刮风', 'guā fēng', 'windy'),
        ('下雨', 'xià yǔ', 'raining'),
        ('太阳', 'tài yáng', 'sun'),
        ('月亮', 'yuè liang', 'moon'),
        ('星星', 'xīng xīng', 'stars'),
        ('大树', 'dà shù', 'tree'),
        ('花朵', 'huā duǒ', 'flower'),
    ]
)

# ── 9. 常用句型 ─────────────────────────────────────────
add_section_heading('九、常用句型（写作必备）')
add_vocab_table(
    ['句型', '英文用法'],
    [
        ('虽然……但是……', 'Although... but...'),
        ('因为……所以……', 'Because... therefore...'),
        ('不但……而且……', 'Not only... but also...'),
        ('一边……一边……', 'While doing... also doing...'),
        ('如果……就……', 'If... then...'),
        ('只要……就……', 'As long as... then...'),
        ('终于', 'finally'),
        ('突然', 'suddenly'),
        ('渐渐', 'gradually'),
        ('立刻', 'immediately'),
    ]
)

# ── 10. 常用成语 ────────────────────────────────────────
add_section_heading('十、常用成语')
add_vocab_table(
    ['成语', '拼音', '意思'],
    [
        ('勤能补拙', 'qín néng bǔ zhuō', 'Hard work makes up for lack of talent'),
        ('助人为乐', 'zhù rén wéi lè', 'Joy in helping others'),
        ('废寝忘食', 'fèi qǐn wàng shí', 'So engrossed one forgets to eat/sleep'),
        ('马到成功', 'mǎ dào chéng gōng', 'Immediate success'),
        ('一石二鸟', 'yī shí èr niǎo', 'Kill two birds with one stone'),
        ('半途而废', 'bàn tú ér fèi', 'Give up halfway'),
        ('自食其力', 'zì shí qí lì', 'Be self-reliant'),
        ('精益求精', 'jīng yì qiú jīng', 'Always strive for improvement'),
    ]
)

# ── Tips ────────────────────────────────────────────────
doc.add_paragraph()
add_tip_box([
    '每天学习 10–15 个词语',
    '用词语造句，加深记忆',
    '用听写方式每周复习',
    '阅读短文，找出生词并查字典',
])

# Save
output_path = '/home/user/MicrobiomeStatPlot/小六华文常用词汇手册.docx'
doc.save(output_path)
print(f'Saved: {output_path}')

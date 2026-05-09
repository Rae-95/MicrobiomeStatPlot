from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

def set_font(run, bold=False, size=12, color=None, italic=False):
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.bold = bold
    run.font.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_para(text='', bold=False, size=12, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size, color=color, italic=italic)
    return p

def add_section_title(num, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1F497D')
    pPr.append(shd)
    p.paragraph_format.left_indent = Cm(0.3)
    run = p.add_run(f'第{num}题  {title}')
    set_font(run, bold=True, size=13, color=(0xFF, 0xFF, 0xFF))
    return p

def add_divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run('─' * 48)
    set_font(run, size=9, color=(0xBB, 0xBB, 0xBB))

def add_blank_line(n, label='', size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f'{label}{"＿" * n}')
    set_font(run, size=size)

def add_answer_line(prefix='答：', n=30):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f'{prefix}{"＿" * n}')
    set_font(run, size=11)

# ── HEADER ─────────────────────────────────────────────
add_para('新加坡中一华文词汇运用测验', bold=True, size=18,
         color=(0x1F, 0x49, 0x7D), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2)
add_para('Secondary 1 Chinese Vocabulary Test', bold=False, size=11,
         color=(0x70, 0x70, 0x70), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

# Info table
info_table = doc.add_table(rows=1, cols=3)
info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cells = info_table.rows[0].cells
for cell, label in zip(cells, ['姓名：＿＿＿＿＿＿＿＿', '班级：＿＿＿＿＿＿＿＿', '日期：＿＿＿＿＿＿＿＿']):
    cell.text = label
    for run in cell.paragraphs[0].runs:
        set_font(run, size=11)

doc.add_paragraph()

# Word bank box
add_para('【词语表】', bold=True, size=12, color=(0x1F, 0x49, 0x7D), space_after=2)
wb_table = doc.add_table(rows=4, cols=5)
wb_table.style = 'Table Grid'
wb_table.alignment = WD_TABLE_ALIGNMENT.CENTER
words = ['设施','技能','别扭','冤枉','锻炼',
         '达人','科技','口头禅','小事一桩','人缘',
         '推荐','惹麻烦','红枣','大吉大利','疼爱',
         '儿媳妇','感激','莫名其妙','疑惑','']
for i, word in enumerate(words):
    row_idx = i // 5
    col_idx = i % 5
    cell = wb_table.cell(row_idx, col_idx)
    cell.text = word
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cell.paragraphs[0].runs:
        set_font(run, bold=True, size=12, color=(0x1F, 0x49, 0x7D))
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'DEEAF1')
    tcPr.append(shd)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 第一题：配对题
# ══════════════════════════════════════════════════════════
add_section_title('一', '配对题——将词语与正确的解释配对  （19分）')
add_para('请将左边的词语与右边正确的解释用线连起来，或将字母填入括号中。', size=11, italic=True, space_after=4)

match_table = doc.add_table(rows=19, cols=4)
match_table.alignment = WD_TABLE_ALIGNMENT.CENTER

left_words = [
    '1. 设施', '2. 技能', '3. 别扭', '4. 冤枉',
    '5. 锻炼', '6. 达人', '7. 科技', '8. 口头禅',
    '9. 小事一桩', '10. 人缘', '11. 推荐', '12. 惹麻烦',
    '13. 红枣', '14. 大吉大利', '15. 疼爱', '16. 儿媳妇',
    '17. 感激', '18. 莫名其妙', '19. 疑惑',
]
right_defs = [
    'A. 儿子的妻子',
    'B. 祝福吉祥如意的话',
    'C. 经常挂在嘴边的话',
    'D. 对别人的帮助深感谢意',
    'E. 硬说人做了没做的事，造成不公正',
    'F. 为人处事受欢迎的程度',
    'G. 科学与技术的总称',
    'H. 感到奇怪，说不出原因',
    'I. 一种滋补果实，常用于煲汤',
    'J. 极其关爱',
    'K. 通过练习使身体或某方面能力提高',
    'L. 心里不舒服，感到不自然',
    'M. 为别人介绍，建议接受某事物',
    'N. 某方面非常出色的人',
    'O. 非常简单的小事',
    'P. 引起纠纷或困难',
    'Q. 为完成某项任务所具备的技术和能力',
    'R. 心中不明白，感到困惑',
    'S. 为生活提供方便的建筑或器械设备',
]

for i in range(19):
    row = match_table.rows[i].cells
    row[0].text = left_words[i]
    row[1].text = '（  ）'
    row[2].text = ''
    row[3].text = right_defs[i]
    for run in row[0].paragraphs[0].runs:
        set_font(run, bold=True, size=11)
    for run in row[1].paragraphs[0].runs:
        set_font(run, size=11)
    for run in row[3].paragraphs[0].runs:
        set_font(run, size=11)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 第二题：填空题
# ══════════════════════════════════════════════════════════
add_section_title('二', '填空题——从词语表中选出最合适的词语填入空格  （10分）')
add_para('每题只用一个词语，每个词语只可使用一次。', size=11, italic=True, space_after=4)

fill_questions = [
    ('1.', '小明非常努力，每天早起', '＿＿＿＿', '，身体越来越健壮。'),
    ('2.', '他待人友善，在班上', '＿＿＿＿', '极好，大家都喜欢和他在一起。'),
    ('3.', '奶奶煲了一锅加了', '＿＿＿＿', '的鸡汤，香气四溢。'),
    ('4.', '老师向同学们', '＿＿＿＿', '了一本非常有趣的课外读物。'),
    ('5.', '他不小心打翻了水杯，', '＿＿＿＿', '，弄湿了同学的作业本。'),
    ('6.', '妈妈对小妹妹非常', '＿＿＿＿', '，什么好东西都留给她。'),
    ('7.', '新年期间，大家互相祝福"', '＿＿＿＿', '"，气氛非常热闹。'),
    ('8.', '他是学校的电脑', '＿＿＿＿', '，什么难题都难不倒他。'),
    ('9.', '她', '＿＿＿＿', '地感到有人在注视她，却四处张望也看不见任何人。'),
    ('10.', '对于老师的帮助，她内心非常', '＿＿＿＿', '，特地写了一封感谢信。'),
]

for q in fill_questions:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r0 = p.add_run(q[0] + ' ' + q[1])
    set_font(r0, size=11)
    r1 = p.add_run(q[2])
    set_font(r1, size=11, color=(0x1F, 0x49, 0x7D), bold=True)
    r2 = p.add_run(q[3])
    set_font(r2, size=11)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 第三题：选择题
# ══════════════════════════════════════════════════════════
add_section_title('三', '选择题——选出词语运用正确的句子  （5分）')
add_para('请圈出正确答案。', size=11, italic=True, space_after=4)

mc_questions = [
    ('1. "别扭"的正确用法是：',
     'A. 他和同桌闹了矛盾，两人相处得很别扭。',
     'B. 这道数学题太别扭了，所以很简单。',
     'C. 今天天气别扭，阳光灿烂。',
     'D. 她的歌声别扭，大家都很喜欢听。', 'A'),
    ('2. "莫名其妙"的正确用法是：',
     'A. 他学习莫名其妙，成绩一直很好。',
     'B. 她莫名其妙地哭了起来，让大家不知所措。',
     'C. 这道菜莫名其妙，非常美味。',
     'D. 操场莫名其妙，同学们都在玩耍。', 'B'),
    ('3. "冤枉"的正确用法是：',
     'A. 他说这本书冤枉，值得一读。',
     'B. 那个地方冤枉，风景很美。',
     'C. 他没有偷东西，却被人冤枉，心里很难过。',
     'D. 她冤枉地完成了作业。', 'C'),
    ('4. "小事一桩"的正确用法是：',
     'A. 帮你搬这几本书，小事一桩，不用客气。',
     'B. 这次考试小事一桩，大家都很紧张。',
     'C. 小事一桩的风景让人心旷神怡。',
     'D. 他小事一桩地走进了教室。', 'A'),
    ('5. "口头禅"的正确用法是：',
     'A. 他的口头禅是"没问题"，不管什么事都说没问题。',
     'B. 她每天念口头禅来保佑自己。',
     'C. 这首歌的口头禅非常好听。',
     'D. 老师在黑板上写下了口头禅。', 'A'),
]

for q_num, (question, a, b, c, d, ans) in enumerate(mc_questions):
    add_para(question, bold=True, size=11, space_after=1)
    for opt in [a, b, c, d]:
        add_para('    ' + opt, size=11, space_after=1)
    add_para(f'    （正确答案：{ans}）', size=10, color=(0xBB, 0xBB, 0xBB), italic=True, space_after=6)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 第四题：造句题
# ══════════════════════════════════════════════════════════
add_section_title('四', '造句题——用所给词语造句  （20分）')
add_para('请用以下词语各造一个完整、通顺的句子，不少于15个字。', size=11, italic=True, space_after=4)

sentence_words = ['技能', '科技', '疑惑', '儿媳妇', '惹麻烦']
for i, word in enumerate(sentence_words):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f'{i+1}. 【{word}】')
    set_font(r, bold=True, size=12, color=(0x1F, 0x49, 0x7D))

    add_answer_line('    答：', 28)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 第五题：改错题
# ══════════════════════════════════════════════════════════
add_section_title('五', '改错题——找出句中运用错误的词语并改正  （10分）')
add_para('下列句子各有一个词语运用不当，请找出并在横线上写出正确的词语。', size=11, italic=True, space_after=4)

wrong_sentences = [
    ('1.', '他经常疼爱同学，大家都不太喜欢和他玩。',
     '（"疼爱"应改为"__________"）'),
    ('2.', '这道数学题太感激了，我一下子就算出来了。',
     '（"感激"应改为"__________"）'),
    ('3.', '妈妈推荐我做错了事，让我很难过。',
     '（"推荐"应改为"__________"）'),
    ('4.', '运动会上，同学们展示了自己的口头禅，赢得了热烈掌声。',
     '（"口头禅"应改为"__________"）'),
    ('5.', '新年时，大家互相说"疑惑"，表达对新年的祝福。',
     '（"疑惑"应改为"__________"）'),
]

for num, sentence, correction in wrong_sentences:
    add_para(f'{num} {sentence}', size=11, space_after=1)
    add_para(f'    {correction}', size=11, color=(0x70, 0x70, 0x70), space_after=8)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════
# 第六题：阅读短文填空
# ══════════════════════════════════════════════════════════
add_section_title('六', '阅读理解填空——读短文，从词语表中选词填入空格  （10分）')
add_para('阅读下面的短文，选用词语表中的词语填入括号内。', size=11, italic=True, space_after=4)

passage = (
    '    阿强是班上出了名的电脑（    ），他对（    ）的运用非常熟练，这是他最大的（    ）。'
    '同学们遇到电脑问题，都会来（    ）他。每次他都笑着说："（    ），包在我身上！"'
    '因为他待人真诚，所以（    ）很好，大家都很喜欢他。\n'
    '    有一次，他帮同学修电脑时，不小心删除了一份重要文件，（    ）了。'
    '同学起初感到（    ），不知道发生了什么事，后来看到消失的文件，才明白是怎么回事。'
    '阿强感到非常愧疚，连忙道歉。同学没有（    ）他，反而（    ）他帮自己想办法恢复文件。'
    '最终，文件成功找回，两人的友谊也更深了。'
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(8)
run = p.add_run(passage)
set_font(run, size=11)

add_para('（参考词语：达人、科技、技能、推荐、小事一桩、人缘、惹麻烦、疑惑、冤枉、感激）',
         size=10, color=(0x70, 0x70, 0x70), italic=True, space_after=4)

# Score summary
doc.add_paragraph()
add_divider()
score_table = doc.add_table(rows=2, cols=7)
score_table.style = 'Table Grid'
score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['题目', '第一题', '第二题', '第三题', '第四题', '第五题', '第六题']
scores  = ['满分', '19分', '10分', '5分', '20分', '10分', '10分']
for i, (h, s) in enumerate(zip(headers, scores)):
    c1 = score_table.cell(0, i)
    c2 = score_table.cell(1, i)
    c1.text = h
    c2.text = s
    for row_cell, is_header in [(c1, True), (c2, False)]:
        for run in row_cell.paragraphs[0].runs:
            set_font(run, bold=is_header, size=11)
        row_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tc = row_cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1F497D' if is_header else 'DEEAF1')
        tcPr.append(shd)
        if is_header:
            for run in row_cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

doc.add_paragraph()
add_para('总分：________／74分        得分：________', bold=True, size=12,
         align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)

output_path = '/home/user/MicrobiomeStatPlot/中一华文词汇测验.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
